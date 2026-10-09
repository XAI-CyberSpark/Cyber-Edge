"""Read tool-definition worksheets locally; never execute spreadsheet content."""
import ast
import io
import json
import re
import zipfile
from openpyxl import load_workbook


def _category(row):
    from app.services.capability_categories import classify
    return classify(str(row.get('name') or ''), description=str(row.get('description') or ''),
                    domain=str(row.get('一级skill') or ''), group=str(row.get('二级分类') or ''),
                    title=str(row.get('title') or ''))


def read_tools(filename, data):
    from app.services.capability_library import ImportProblem
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        if len(archive.infolist()) > 1000 or sum(i.file_size for i in archive.infolist()) > 30 * 1024 * 1024:
            raise ImportProblem('Excel 解压后过大，请拆分工作表后重试。')
    book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    warnings, records, seen = [], [], {}
    inspected = 0
    try:
        sheets = sorted(book.worksheets, key=lambda sheet: (sheet.title != '二级分类明细', sheet.title != 'Sheet1'))
        tables = 0
        for sheet in sheets:
            headers = None
            for number, cells in enumerate(sheet.iter_rows(), 1):
                inspected += 1
                if inspected > 25000:
                    raise ImportProblem('Excel 行数过多，请拆分后导入。')
                values = [cell.value for cell in cells]
                if headers is None:
                    if number > 10:
                        break
                    possible = [str(value).strip() if value is not None else '' for value in values]
                    if {'name', 'description', 'inputSchema'} <= set(possible):
                        headers = possible
                        tables += 1
                    continue
                if not any(value is not None for value in values):
                    continue
                row = dict(zip(headers, values))
                name = str(row.get('name') or '').strip()
                if not name:
                    warnings.append(f'{sheet.title} 第{number}行缺少接口名称，未导入。')
                    continue
                if any(getattr(cell, 'data_type', '') == 'f' for cell in cells):
                    raise ImportProblem(f'{sheet.title} 第{number}行包含公式，请先将接口资料转换为值。')
                raw_schema = row.get('inputSchema')
                identity = (str(row.get('description') or ''), str(raw_schema))
                if name in seen:
                    if seen[name] != identity:
                        warnings.append(f'{sheet.title} 第{number}行：同名接口 {name} 内容不同，保留优先明细表记录，请核对。')
                    continue
                seen[name] = identity
                missing = []
                try:
                    if isinstance(raw_schema, str):
                        try:
                            schema = json.loads(raw_schema)
                        except json.JSONDecodeError:
                            schema = ast.literal_eval(raw_schema)
                    else:
                        schema = raw_schema
                    if not isinstance(schema, dict) or not isinstance(schema.get('properties', {}), dict) or not isinstance(schema.get('required', []), list):
                        raise ValueError('schema structure')
                    json.dumps(schema)
                    for key in schema.get('required', []):
                        prop = schema.get('properties', {}).get(key)
                        if not isinstance(prop, dict):
                            missing.append(f'参数 {key} 的定义待补充')
                        elif any(word in str(prop.get('description', '')) for word in ('不传', '非必填', '可选参数')):
                            missing.append(f'参数 {key} 必填约束与说明冲突，待核对')
                except (ValueError, SyntaxError, TypeError, RecursionError):
                    schema = str(raw_schema or '')
                    missing.append('输入参数格式待核对')
                description = str(row.get('description') or '')
                output = row.get('outputSchema') or row.get('outputs') or ''
                if not output:
                    parts = re.split(r'[:：]', description, maxsplit=1)
                    if len(parts) == 2 and parts[1].strip():
                        output = '原文返回说明（未验证）：' + parts[1].strip()
                label = re.split(r'[:：,，。\n]', description, maxsplit=1)[0].strip()
                if len(label) > 44:
                    label = str(row.get('title') or name)
                records.append({
                    'name': label or str(row.get('title') or name), 'interface': name,
                    'description': description, 'inputs': schema, 'outputs': output,
                    'permissions': row.get('permissions'), 'conditions': row.get('conditions'),
                    'platform': row.get('platform'), 'version': row.get('version'),
                    'category': _category(row), 'source_group': str(row.get('一级业务域') or row.get('一级skill') or ''),
                    'source_category': str(row.get('二级分类') or ''),
                    'source': f'{filename} / {sheet.title}!A{number}:H{number}',
                    'evidence': json.dumps(row, ensure_ascii=False, default=str), 'missing': missing,
                })
        if not tables:
            raise ImportProblem('没有找到接口明细表。Excel 需包含 name、description、inputSchema 列，可附分类、title 和返回说明。')
        if not records:
            raise ImportProblem('接口明细表没有有效记录。')
        if len(records) > 5000:
            raise ImportProblem('接口超过5000项，请分批导入。')
        return records, warnings
    finally:
        book.close()
